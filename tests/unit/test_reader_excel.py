"""Excel and ODS: title rows, merged headers, native cell types, sheet selection."""

from __future__ import annotations

import datetime as dt

import openpyxl
import pyarrow as pa
import pytest

import aimart_testkit as kit
from aimart_parquet.engine.limits import Limits
from aimart_parquet.engine.readers.excel import (
    _column_array,
    _fill_merged_headers,
    _header_row,
    _looks_like_header,
    _stringify,
)
from aimart_parquet.errors import InputTooLarge, InvalidInput, InvalidRequest


@pytest.fixture
def workbook(tmp_path):
    return kit.messy_workbook(tmp_path / "book" / "sales.xlsx")


class TestMessyWorkbook:
    def test_first_sheet_with_title_rows_and_merged_header(self, workbook) -> None:
        result = kit.convert_path(workbook)
        report = result.report
        assert report.source_format == "xlsx"
        assert report.sheet == "Report"
        assert report.available_sheets == ["Report", "Empty", "Numbers"]
        assert "skipped 2 title or blank rows above the table" in report.warnings
        assert "dropped 1 empty rows" in report.warnings
        assert [c.name for c in report.columns] == [
            "region",
            "sales",
            "sales_2",  # the merged "Sales" header spans two columns
            "date",
            "mixed",
            "flag",
            "time",
        ]
        rows = result.rows
        assert rows["region"] == ["North", "South", "East"]
        assert rows["sales"] == [10, 20, 30]
        assert rows["sales_2"] == [1.5, 2.5, 3.25]
        assert rows["date"] == [
            dt.datetime(2025, 1, 1),
            dt.datetime(2025, 1, 2),
            dt.datetime(2025, 1, 3, 12, 0),
        ]
        assert rows["mixed"] == [5.0, None, 7.5]  # "n/a" is a null token in a typed column
        assert rows["flag"] == [True, False, True]
        assert rows["time"] == [dt.time(10, 30), dt.time(11, 0), dt.time(12, 0)]
        assert result.column("sales").type == "int64"
        assert result.column("mixed").inferred_from_text

    @pytest.mark.parametrize("sheet", ["Numbers", 2])
    def test_sheet_by_name_or_index(self, workbook, sheet) -> None:
        result = kit.convert_path(workbook, sheet=sheet)
        assert result.report.sheet == "Numbers"
        # A first row of numbers is data, not a header.
        assert result.rows == {"column_1": [1, 4], "column_2": [2, 5], "column_3": [3, 6]}

    def test_forcing_a_header(self, workbook) -> None:
        result = kit.convert_path(workbook, sheet="Numbers", header=True)
        assert result.rows == {"c_1": [4], "c_2": [5], "c_3": [6]}

    def test_unknown_sheet_lists_the_alternatives(self, workbook) -> None:
        with pytest.raises(InvalidRequest) as info:
            kit.convert_path(workbook, sheet="Nope")
        assert info.value.details == {"available_sheets": ["Report", "Empty", "Numbers"]}

    def test_sheet_index_out_of_range(self, workbook) -> None:
        with pytest.raises(InvalidRequest, match="out of range"):
            kit.convert_path(workbook, sheet=7)

    def test_empty_sheet_by_name(self, workbook) -> None:
        with pytest.raises(InvalidInput, match="empty"):
            kit.convert_path(workbook, sheet="Empty")

    def test_first_non_empty_sheet_is_the_default(self, tmp_path) -> None:
        wb = openpyxl.Workbook()
        first = wb.active
        assert first is not None
        first.title = "Blank"
        wb.create_sheet("Data").append(["a", "b"])
        wb["Data"].append([1, 2])
        path = tmp_path / "b.xlsx"
        wb.save(path)
        result = kit.convert_path(path)
        assert result.report.sheet == "Data"
        assert result.rows == {"a": [1], "b": [2]}

    def test_all_sheets_empty(self, tmp_path) -> None:
        path = tmp_path / "empty.xlsx"
        openpyxl.Workbook().save(path)
        with pytest.raises(InvalidInput, match="empty"):
            kit.convert_path(path)

    def test_corrupt_workbook(self, tmp_path) -> None:
        path = kit.write(
            tmp_path / "bad.xlsx", kit.zip_bytes({"[Content_Types].xml": b"<x", "xl/a": b""})
        )
        with pytest.raises(InvalidInput, match="Could not read the spreadsheet"):
            kit.convert_path(path)

    def test_ods_is_detected_and_read(self, tmp_path) -> None:
        path = kit.minimal_ods(tmp_path / "sheet.bin", [["name", "qty"], ["a", 3], ["b", 4.5]])
        result = kit.convert_path(path, filename=None)  # no name: magic bytes only
        assert result.report.source_format == "ods"
        assert result.rows == {"name": ["a", "b"], "qty": [3.0, 4.5]}

    def test_xlsx_zip_bomb_is_checked_before_parsing(self, workbook) -> None:
        with pytest.raises(InputTooLarge):
            kit.convert_path(workbook, limits=Limits(max_decompressed_bytes=100))


class TestHelpers:
    def test_header_row_skips_sparse_title_rows(self) -> None:
        rows = [["Title", None, None, None], [None] * 4, ["a", "b", "c", "d"], [1, 2, 3, 4]]
        assert _header_row(rows, 4) == 2

    def test_header_row_for_narrow_tables(self) -> None:
        assert _header_row([[None, None], ["a", None]], 2) == 1

    def test_header_row_falls_back_to_the_first_non_empty_row(self) -> None:
        rows = [[None] * 6] * 31 + [["x", None, None, None, None, None]]
        assert _header_row(rows, 6) == 31

    def test_looks_like_header(self) -> None:
        assert _looks_like_header(["a", None, " "])
        assert not _looks_like_header(["a", 1])

    @pytest.mark.parametrize(
        ("value", "expected"),
        [
            (None, None),
            ("  ", None),
            (True, "true"),
            (3.0, "3"),
            (3.5, "3.5"),
            (2.0**60, str(2.0**60)),
            (dt.date(2024, 1, 2), "2024-01-02"),
            (" x ", "x"),
        ],
    )
    def test_stringify(self, value, expected) -> None:
        assert _stringify(value) == expected

    @pytest.mark.parametrize(
        ("values", "arrow_type", "infer"),
        [
            ([None, " "], pa.string(), False),
            ([True, False, None], pa.bool_(), False),
            ([1, 2.0, None], pa.int64(), False),
            ([1, 2.5], pa.float64(), False),
            ([2.0**60, 1], pa.float64(), False),
            ([dt.date(2024, 1, 1), None], pa.date32(), False),
            ([dt.date(2024, 1, 1), dt.datetime(2024, 1, 2, 3)], pa.timestamp("us"), False),
            ([dt.time(1, 2)], pa.time64("us"), False),
            ([1, "x"], pa.string(), True),
        ],
    )
    def test_column_array_keeps_native_types_when_cells_agree(
        self, values, arrow_type, infer
    ) -> None:
        array, needs_inference = _column_array(values)
        assert array.type == arrow_type
        assert needs_inference is infer

    def test_fill_merged_headers_respects_the_sheet_offset(self) -> None:
        header = ["A", None, None, "D"]
        # Sheet data starts at row 2, column 1; the merge covers columns 1-3 of row 2.
        _fill_merged_headers(header, [((2, 1), (2, 3))], 0, (2, 1))
        assert header == ["A", "A", "A", "D"]

    def test_fill_merged_headers_ignores_other_rows_and_keeps_text(self) -> None:
        header = ["A", "B", None]
        _fill_merged_headers(header, [((5, 0), (5, 2)), ((0, 0), (0, 2))], 0, (0, 0))
        assert header == ["A", "B", "A"]
        _fill_merged_headers(header, [], 0, (0, 0))
        assert header == ["A", "B", "A"]
