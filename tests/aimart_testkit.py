"""Shared test helpers: a recording fake facilitator, app harnesses, wire decoders, file makers.

Imported by test modules as ``aimart_testkit`` (``tests/conftest.py`` puts this directory on
``sys.path``; the suite runs with ``--import-mode=importlib``, which does not). Fixtures live in
conftest.py; anything a test needs to *call* with its own arguments lives here.
"""

from __future__ import annotations

import base64
import contextlib
import io
import json
import zipfile
from dataclasses import dataclass, field
from typing import TYPE_CHECKING, Any

import httpx
from x402 import SettleResponse, SupportedKind, SupportedResponse, VerifyResponse

if TYPE_CHECKING:
    from collections.abc import AsyncIterator, Iterable
    from pathlib import Path

    import mcp
    from fastapi import FastAPI
    from x402 import x402Client

    from aimart_parquet.config import Settings
    from aimart_parquet.payments import Payments

# Receive-only test addresses. Distinct values make "which address ended up where" assertions
# meaningful: payTo must be ours, never the payer's.
PAY_TO_EVM = "0x" + "ab" * 20
PAY_TO_SVM = "9xQeWvG816bUx9EPjHmaT23yvVM2ZWbrrpZb9PusVFin"
PAYER = "0x" + "cd" * 20
FEE_PAYER = "11111111111111111111111111111111"
TX_HASH = "0x" + "ef" * 32
PAYER_KEY = "0x" + "11" * 32  # deterministic client wallet; never funded anywhere

BASE_SEPOLIA = "eip155:84532"
BASE_MAINNET = "eip155:8453"
USDC_BASE_SEPOLIA = "0x036CbD53842c5426634e7929541eC2318f3dCF7e"
USDC_BASE_MAINNET = "0x833589fCD6eDb6E08f4c7C32D4f71b54bdA02913"
SOLANA_DEVNET = "solana:EtWTRABZaYq6iMfeYKouRu166VU2xqa1"
USDC_SOLANA_DEVNET = "4zMMC9srt5Ri5X14GAgXhaHii3GnPAEERYPJgZJDncDU"

BASE_URL = "http://test"


# ---------------------------------------------------------------------------------- facilitator
@dataclass
class FakeFacilitator:
    """An in-memory facilitator that records every verify/settle call.

    The x402 SDK calls ``get_supported`` synchronously during initialization and ``verify`` /
    ``settle`` asynchronously. Flip the attributes between calls to script failures.
    """

    networks: tuple[str, ...] = (BASE_SEPOLIA,)
    kind_extra: dict[str, dict[str, Any]] = field(default_factory=dict)
    valid: bool = True
    settle_ok: bool = True
    verify_error: Exception | None = None
    settle_error: Exception | None = None
    supported_error: Exception | None = None
    verified: list[tuple[Any, Any]] = field(default_factory=list)
    settled: list[tuple[Any, Any]] = field(default_factory=list)
    supported_calls: int = 0

    def get_supported(self) -> SupportedResponse:
        self.supported_calls += 1
        if self.supported_error is not None:
            raise self.supported_error
        extra = dict(self.kind_extra)
        for network in self.networks:
            if network.startswith("solana:"):
                extra.setdefault(network, {"feePayer": FEE_PAYER})
        return SupportedResponse(
            kinds=[
                SupportedKind(x402_version=2, scheme="exact", network=n, extra=extra.get(n))
                for n in self.networks
            ]
        )

    async def verify(self, payload: Any, requirements: Any) -> VerifyResponse:
        self.verified.append((payload, requirements))
        if self.verify_error is not None:
            raise self.verify_error
        if not self.valid:
            return VerifyResponse(is_valid=False, invalid_reason="insufficient_funds")
        return VerifyResponse(is_valid=True, payer=PAYER)

    async def settle(self, payload: Any, requirements: Any) -> SettleResponse:
        self.settled.append((payload, requirements))
        if self.settle_error is not None:
            raise self.settle_error
        if not self.settle_ok:
            return SettleResponse(
                success=False,
                error_reason="insufficient_funds",
                transaction="",
                network=requirements.network,
            )
        return SettleResponse(
            success=True,
            transaction=TX_HASH,
            network=requirements.network,
            payer=PAYER,
            amount=requirements.amount,
        )


# ---------------------------------------------------------------------------------- wire
def decode_b64_json(value: str) -> dict[str, Any]:
    """Decode a base64 JSON header exactly as a client would — no SDK helpers in between."""
    decoded: dict[str, Any] = json.loads(base64.b64decode(value))
    return decoded


def payment_required(response: httpx.Response) -> dict[str, Any]:
    assert response.status_code == 402, response.text
    header = response.headers.get("payment-required")
    assert header, f"402 without a PAYMENT-REQUIRED header: {dict(response.headers)}"
    return decode_b64_json(header)


def payment_response(response: httpx.Response) -> dict[str, Any]:
    header = response.headers.get("payment-response")
    assert header, f"no PAYMENT-RESPONSE header: {dict(response.headers)}"
    return decode_b64_json(header)


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def paying_client(wallet_key: str = PAYER_KEY) -> x402Client:
    import eth_account
    from x402 import x402Client
    from x402.mechanisms.evm import EthAccountSigner
    from x402.mechanisms.evm.exact import ExactEvmClientScheme

    account = eth_account.Account.from_key(wallet_key)
    client = x402Client()
    client.register(BASE_SEPOLIA, ExactEvmClientScheme(EthAccountSigner(account)))
    return client


# ---------------------------------------------------------------------------------- apps
@dataclass
class AppHarness:
    app: FastAPI
    settings: Settings
    facilitator: FakeFacilitator | None
    payments: Payments | None
    http: httpx.AsyncClient  # never pays: what an unpaid agent sees
    paying: httpx.AsyncClient  # signs and retries on 402, like a real x402 agent


@contextlib.asynccontextmanager
async def running_app(
    settings: Settings,
    *,
    facilitator: FakeFacilitator | None = None,
    paid: bool = True,
    client: x402Client | None = None,
) -> AsyncIterator[AppHarness]:
    """Build the real ASGI app, run its lifespan, and hand back unpaid + paying clients."""
    from x402.http.clients.httpx import x402AsyncTransport

    from aimart_parquet.api.app import create_app
    from aimart_parquet.payments import Payments

    payments: Payments | None = None
    if paid:
        facilitator = facilitator or FakeFacilitator()
        payments = Payments(settings, [facilitator])
        app = create_app(settings, payments=payments)
    else:
        app = create_app(settings, payments=False)
    async with app.router.lifespan_context(app):
        plain = httpx.ASGITransport(app=app)
        paying = x402AsyncTransport(client or paying_client(), httpx.ASGITransport(app=app))
        async with (
            httpx.AsyncClient(transport=plain, base_url=BASE_URL) as http,
            httpx.AsyncClient(transport=paying, base_url=BASE_URL) as paying_http,
        ):
            yield AppHarness(app, settings, facilitator, payments, http, paying_http)


@contextlib.asynccontextmanager
async def mcp_over_http(app: FastAPI) -> AsyncIterator[mcp.Client]:
    """An MCP client speaking Streamable HTTP to the mounted /mcp endpoint, in-process."""
    import httpx2
    import mcp
    from mcp.client.streamable_http import streamable_http_client

    async with httpx2.AsyncClient(
        transport=httpx2.ASGITransport(app=app), base_url=BASE_URL
    ) as http_client:
        transport = streamable_http_client(f"{BASE_URL}/mcp", http_client=http_client)
        async with mcp.Client(transport) as client:
            yield client


async def mcp_payment_meta(client: x402Client, structured: dict[str, Any]) -> dict[str, Any]:
    """Turn an MCP PaymentRequired result into the params._meta a paying client sends."""
    from x402 import parse_payment_required

    payload = await client.create_payment_payload(parse_payment_required(structured))
    return {"x402/payment": payload.model_dump(by_alias=True, exclude_none=True)}


# ---------------------------------------------------------------------------------- engine
@dataclass
class Converted:
    report: Any  # ConversionReport
    path: Path

    @property
    def table(self) -> Any:  # pyarrow.Table
        import pyarrow.parquet as pq

        return pq.read_table(self.path)

    @property
    def rows(self) -> dict[str, list[Any]]:
        return dict(self.table.to_pydict())

    def column(self, name: str) -> Any:  # ColumnReport
        return next(c for c in self.report.columns if c.name == name)


_counter = iter(range(10**9))


def convert_path(
    src: Path, *, filename: str | None = "auto", limits: Any = None, **options: Any
) -> Converted:
    """Run the engine on a file in a fresh work directory next to it."""
    from aimart_parquet.engine.convert import convert_file
    from aimart_parquet.engine.limits import Limits
    from aimart_parquet.engine.models import ConvertOptions

    workdir = src.parent / f"work{next(_counter)}"
    workdir.mkdir()
    out = workdir / "out.parquet"
    report = convert_file(
        src,
        out,
        filename=src.name if filename == "auto" else filename,
        options=ConvertOptions(**options),
        workdir=workdir,
        limits=limits or Limits(),
    )
    return Converted(report, out)


def convert_bytes(
    tmp_path: Path, name: str, data: bytes | str, *, limits: Any = None, **options: Any
) -> Converted:
    src = write(tmp_path / f"in{next(_counter)}" / name, data)
    return convert_path(src, limits=limits, **options)


# ---------------------------------------------------------------------------------- files
MESSY_CSV = (
    "﻿Customer ID,Customer Name,Amount (EUR),Signup Date,Active,,Customer Name\n"
    '00123,José,"1.234,56",03/04/2025,yes,x,dup\n'
    '00124,Zoë,"12,5",25/12/2024,no,,dup\n'
    "00125,NA,7,13/01/2025,yes,y\n"
)


def write(path: Path, data: bytes | str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, str):
        path.write_text(data, encoding="utf-8")
    else:
        path.write_bytes(data)
    return path


def zip_bytes(members: dict[str, bytes], *, compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return buf.getvalue()


def csv_text(header: Iterable[str], rows: Iterable[Iterable[object]]) -> str:
    lines = [",".join(header)]
    lines += [",".join("" if v is None else str(v) for v in row) for row in rows]
    return "\n".join(lines) + "\n"


def write_parquet(path: Path, columns: dict[str, list[Any]], **kwargs: Any) -> Path:
    import pyarrow as pa
    import pyarrow.parquet as pq

    path.parent.mkdir(parents=True, exist_ok=True)
    pq.write_table(pa.table(columns), path, **kwargs)
    return path


def messy_workbook(path: Path) -> Path:
    """A workbook like real exports: a title in A1, a blank row, a merged header, mixed cells."""
    import datetime as dt

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    ws["A1"] = "Quarterly sales report"
    ws.append([])
    ws.append(["Region", "Sales", None, "Date", "Mixed", "Flag", "Time"])
    ws.merge_cells("B3:C3")
    ws.append(["North", 10, 1.5, dt.date(2025, 1, 1), 5, True, dt.time(10, 30)])
    ws.append(["South", 20, 2.5, dt.date(2025, 1, 2), "n/a", False, dt.time(11, 0)])
    ws.append([None] * 7)
    ws.append(["East", 30, 3.25, dt.datetime(2025, 1, 3, 12, 0), 7.5, True, dt.time(12, 0)])
    wb.create_sheet("Empty")
    numbers = wb.create_sheet("Numbers")
    numbers.append([1, 2, 3])
    numbers.append([4, 5, 6])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def minimal_ods(path: Path, rows: list[list[str | float]]) -> Path:
    """A hand-built OpenDocument spreadsheet (no ODS writer is a test dependency)."""
    from xml.sax.saxutils import escape

    def cell(v: str | float) -> str:
        if isinstance(v, str):
            return f'<table:table-cell office:value-type="string"><text:p>{escape(v)}</text:p></table:table-cell>'
        return f'<table:table-cell office:value-type="float" office:value="{v}"><text:p>{v}</text:p></table:table-cell>'

    body = "".join(
        "<table:table-row>" + "".join(cell(v) for v in row) + "</table:table-row>" for row in rows
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f'<office:body><office:spreadsheet><table:table table:name="Sheet1">{body}</table:table>'
        "</office:spreadsheet></office:body></office:document-content>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr("META-INF/manifest.xml", _ODS_MANIFEST, compress_type=zipfile.ZIP_DEFLATED)
    return write(path, buf.getvalue())


_ODS_MANIFEST = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
    'manifest:version="1.2">'
    '<manifest:file-entry manifest:full-path="/" '
    'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
    '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
    "</manifest:manifest>"
)


SURVEY_LABELS = {"gender": "Respondent gender", "score": "Satisfaction score", "name": "Name"}
SURVEY_VALUE_LABELS: dict[str, dict[int | float, str]] = {"gender": {1: "Male", 2: "Female"}}


def survey_file(path: Path, fmt: str) -> Path:
    """An SPSS/Stata/SAS file with variable labels and (where the format allows) value labels."""
    import polars as pl
    import pyreadstat

    df = pl.DataFrame(
        {
            "gender": [1.0, 2.0, 1.0, None],
            "score": [1.5, 2.25, 3.0, 4.0],
            "name": ["a", "b", "c", "d"],
        }
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    labels = dict(SURVEY_LABELS)
    if fmt in {"sav", "zsav"}:
        pyreadstat.write_sav(
            df,
            str(path),
            column_labels=labels,
            variable_value_labels=SURVEY_VALUE_LABELS,
            compress=fmt == "zsav",
        )
    elif fmt == "dta":
        pyreadstat.write_dta(
            df, str(path), column_labels=labels, variable_value_labels=SURVEY_VALUE_LABELS
        )
    elif fmt == "xpt":
        pyreadstat.write_xport(df, str(path), column_labels=labels)
    elif fmt == "por":
        pyreadstat.write_por(df, str(path), column_labels=labels)
    else:  # pragma: no cover - test bug
        raise ValueError(fmt)
    return path
